import io
import sys
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from excel_export import monthly_workbook

class MonthlyWorkbook(unittest.TestCase):
    def sample(self):
        return {'properties':[{'id':'p','name':'Immeuble Test','address':'Conakry','units':2,'type':'Appartements'}], 'tenants':[{'id':'t','name':'=HYPERLINK("https://example.invalid")','phone':'0620000000','property':'p','unit':'A1','rent':2000000,'leaseStart':'2026-01-15','deposit':4000000,'dueDay':5}], 'payments':[{'id':'a','tenant':'t','date':'2026-10-03','month':'2026-09','amount':500000},{'id':'b','tenant':'t','date':'2026-11-02','month':'2026-10','amount':700000}], 'expenses':[{'id':'e','property':'p','date':'2026-10-04','category':'Entretien','label':'Entretien des communs','amount':100000}]}
    def test_totals_dates_and_user_text(self):
        content=monthly_workbook(self.sample(),'2026-10','Partenaire')
        values=load_workbook(io.BytesIO(content),data_only=True)
        formulas=load_workbook(io.BytesIO(content),data_only=False)
        self.assertEqual(values.sheetnames,['Synthese','Loyers','Encaissements','Depenses','Biens','Locataires','Tarifs'])
        for cell,expected in [('B6',2000000),('B7',700000),('B8',1300000),('B10',500000),('B11',100000),('B12',400000)]:self.assertEqual(values['Synthese'][cell].value,expected)
        self.assertEqual(formulas['Loyers']['A7'].data_type,'s')
        self.assertEqual(values['Loyers']['B7'].value,'0620000000')
        self.assertEqual(values['Encaissements']['A7'].value.year,2026)
        self.assertEqual(formulas['Synthese']['B12'].value,'=B10-B11')
    def test_history_legacy_and_empty(self):
        state=self.sample();t=state['tenants'][0];t.pop('phone');t['rent']=3000000;t['rentHistory']=[{'month':'2026-01','amount':2000000},{'month':'2026-11','amount':3000000}]
        result=load_workbook(io.BytesIO(monthly_workbook(state,'2026-10','Test')),data_only=True)
        self.assertEqual(result['Synthese']['B6'].value,2000000)
        empty={k:[] for k in ['properties','tenants','payments','expenses']}
        result=load_workbook(io.BytesIO(monthly_workbook(empty,'2026-10','Vide')),data_only=True)
        self.assertEqual(result['Synthese']['B6'].value,0)
        with self.assertRaises(ValueError):monthly_workbook(empty,'2026-99','Test')
